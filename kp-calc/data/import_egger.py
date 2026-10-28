"""Импорт прайсов Egger (Стандарт-Кубань) в catalog.json: сверка по названию и структуре, отчёт о различиях.
Запуск: python3 -I import_egger.py <ldsp.xlsx> <ps_top.xlsx> <ps_smart.xlsx> <catalog.json> <отчёт.md>
"""
import openpyxl, json, re, sys
ldsp_f, top_f, smart_f, cat_f, rep_f = sys.argv[1:6]

def norm(s):
    s = str(s or '').lower().replace('ё', 'е').replace('\xa0', ' ')
    s = re.sub(r'\(новинка\)|perfectsense top|perfect sense top', '', s)
    return re.sub(r'[^0-9a-zа-я]', '', s)

# ЛДСП Egger: код C, наименование D, структура E, 16 мм — G; строки «ЛМДФ» (столбец B) — не ЛДСП
ws = openpyxl.load_workbook(ldsp_f, data_only=True)['ЛДСП_Egger']
ldsp, lmdf = [], 0
for r in ws.iter_rows(min_row=15, values_only=True):
    code, name, struct, p16 = r[2], r[3], r[4], r[6]
    if not code or not name or not isinstance(p16, (int, float)) or not p16:
        continue
    if r[1] and 'ЛМДФ' in str(r[1]):
        lmdf += 1; continue
    ldsp.append({'code': str(code).strip(), 'name': f'{str(name).strip()} {str(struct or "").strip()}'.strip(), 'price': float(p16)})

def ps(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    out, last = [], None
    for r in ws.iter_rows(min_row=11, max_col=7, values_only=True):
        if not isinstance(r[0], (int, float)) or not r[1]: continue
        if isinstance(r[5], (int, float)) and r[5]: last = float(r[5])   # 18 мм, объединённые ячейки — цена на группу
        out.append({'code': str(r[1]).strip().replace('\xa0', ' '), 'name': re.sub(r'\s+', ' ', str(r[3]).replace('\xa0', ' ')).strip(), 'price': last})
    return out

cat = json.load(open(cat_f, encoding='utf-8'))
G = {g['name']: g for g in cat['groups']}
report = ['# Загрузка прайсов Egger (Стандарт-Кубань, 03.08.2026)', '',
          'Правило: себестоимость м² = цена листа ÷ площадь листа × 1,2 + 150 ₽. ЛДСП — лист 2800×2070 (5,796 м²), 16 мм. Perfect Sense — лист 2800×1220 (3,416 м²), 18 мм.', '']

def apply(gname, new, area, key=lambda x: x['name']):
    g = G[gname]; old = {norm(i['name']): i for i in g['items']}
    up, down, same, added, rows = [], [], 0, [], []
    seen = set()
    for n in new:
        k = norm(key(n)); o = old.get(k)
        if o is None:
            # запасной ключ: код в старом названии
            o = next((i for kk, i in old.items() if norm(n['code']) and norm(n['code']) in kk), None)
        if o is None:
            added.append(n); continue
        seen.add(id(o))
        d = n['price'] - o['price']
        if d > 0.5: up.append((o['name'], o['price'], n['price']))
        elif d < -0.5: down.append((o['name'], o['price'], n['price']))
        else: same += 1
    missing = [i for i in g['items'] if id(i) not in seen]
    g['items'] = [{'id': f"{g['id']}-{k+1}", 'name': f"{n['code']} {n['name']}".strip(), 'price': n['price']} for k, n in enumerate(new)] + \
                 [{'id': f"{g['id']}-x{k+1}", 'name': i['name'], 'price': i['price'], 'status': 'нет в новом прайсе'} for k, i in enumerate(missing)]
    g['rule'], g['unit'], g['sheetArea'] = 'sheet', 'м²', area
    g['source'] = 'Стандарт-Кубань, прайс от 03.08.2026'
    pct = lambda a, b: f'{(b - a) / a * 100:+.1f} %' if a else '—'
    report.extend([f'## {gname}', '', f'В файле: {len(new)} позиций. Совпало с базой: {len(new) - len(added)} (без изменений {same}, подорожали {len(up)}, подешевели {len(down)}). Новые: {len(added)}. Нет в новом прайсе: {len(missing)} (остаются в базе с пометкой).', ''])
    if up or down:
        report.extend(['| Позиция | Было, ₽/лист | Стало, ₽/лист | Изменение |', '|---|---|---|---|'])
        report.extend(f'| {a} | {b:,.0f} | {c:,.0f} | {pct(b, c)} |'.replace(',', ' ') for a, b, c in sorted(up + down, key=lambda x: -(x[2] - x[1]) / x[1])[:40])
        if len(up + down) > 40: report.append(f'| … ещё {len(up + down) - 40} | | | |')
        report.append('')
    if added: report.extend(['Новые: ' + '; '.join(f"{n['code']} {n['name']} — {n['price']:,.0f}".replace(',', ' ') for n in added[:60]) + (' …' if len(added) > 60 else ''), ''])
    if missing: report.extend(['Нет в новом прайсе: ' + '; '.join(i['name'] for i in missing[:60]) + (' …' if len(missing) > 60 else ''), ''])
    return len(up), len(down), len(added), len(missing)

stats = {}
stats['ЛДСП Egger'] = apply('ЛДСП Egger', ldsp, 5.796)
# ЛДСП Egger как фасад — тот же прайс
G['ЛДСП Egger (фасад)']['items'] = [dict(i, id=i['id'].replace(G['ЛДСП Egger']['id'], G['ЛДСП Egger (фасад)']['id'], 1)) for i in G['ЛДСП Egger']['items']]
G['ЛДСП Egger (фасад)'].update(rule='sheet', unit='м²', sheetArea=5.796, source=G['ЛДСП Egger']['source'])
stats['Egger Perfect Sense TOP'] = apply('Egger Perfect Sense TOP', ps(top_f), 3.416)
stats['Egger Perfect Sense Smart'] = apply('Egger Perfect Sense Smart', ps(smart_f), 3.416)
report += ['## Не загружено', '', f'- ЛМДФ из прайса ЛДСП ({lmdf} строк) и файл «ЛМДФ и МДФ Egger»: таких групп в калькуляторе нет — решить владельцу.',
           '- Столбец P прайса ЛДСП (цены ×0,94 от 16 мм) не используется — уточнить, что это.', '- Лист «кромка»: кромка заложена в +150 ₽/м², отдельно не считается.', '']
cat['version'] = cat.get('version', 1) + 1
json.dump(cat, open(cat_f, 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
open(rep_f, 'w', encoding='utf-8').write('\n'.join(report))
print(stats, 'ЛМДФ пропущено:', lmdf)
